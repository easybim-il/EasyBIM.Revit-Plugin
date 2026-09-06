"""Generate an ACC (Autodesk Construction Cloud) Issues Excel import file.

Standalone utility, not a pyRevit button -- run directly with a normal
Python 3 install (needs `openpyxl`; `pandas` is optional, only used if you
pass a DataFrame into create_acc_issues_excel).

CLI usage:
    python acc_issues_export.py
    python acc_issues_export.py --input-json issues.json --output MyIssues.xlsx

issues.json is a JSON array of objects, e.g.:
    [{"title": "Beam clash at Level 2", "status": "Open", "category": "Clash",
      "due_date": "2026-09-30"}]
"""
import argparse
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# (header text, dict key, alignment, min width, max width)
COLUMNS = [
    (u"Title (Required)",                 u"title",                u"left",   20, 45),
    (u"Status",                            u"status",               u"center", 12, 18),
    (u"Category",                          u"category",             u"left",   14, 25),
    (u"Type",                              u"type",                 u"left",   14, 25),
    (u"Description",                       u"description",          u"left",   30, 50),
    (u"Assigned To",                       u"assigned_to",          u"left",   16, 30),
    (u"Location",                          u"location",             u"left",   16, 30),
    (u"Location details",                  u"location_details",     u"left",   20, 40),
    (u"Due Date (YYYY-MM-DD format)",      u"due_date",             u"center", 16, 22),
    (u"Start Date (YYYY-MM-DD format)",    u"start_date",           u"center", 16, 22),
    (u"Root Cause Category",               u"root_cause_category",  u"left",   16, 28),
    (u"Root Cause",                        u"root_cause",           u"left",   16, 28),
]

HEADER_FILL = PatternFill(u"solid", fgColor=u"1F4E79")
HEADER_FONT = Font(color=u"FFFFFF", bold=True)
ZEBRA_FILL = PatternFill(u"solid", fgColor=u"F9FAFB")
GRID_SIDE = Side(style=u"thin", color=u"D9D9D9")
GRID_BORDER = Border(left=GRID_SIDE, right=GRID_SIDE, top=GRID_SIDE, bottom=GRID_SIDE)
HEADER_ROW_HEIGHT = 28
DATA_ROW_HEIGHT = 22


def _rows_from(issues_data):
    """Normalizes list-of-dicts or a pandas DataFrame into a list of plain
    dicts, so the writing loop below never has to care which one it got."""
    if hasattr(issues_data, u"to_dict"):
        return issues_data.to_dict(orient=u"records")
    return list(issues_data or [])


def _cell_value(row, header_text, dict_key):
    """A row built from a DataFrame with the exact header text as its own
    column name (instead of the snake_case dict_key) is honored too, so
    callers aren't forced into one specific naming convention."""
    if dict_key in row:
        return row[dict_key]
    if header_text in row:
        return row[header_text]
    return u""


def create_acc_issues_excel(issues_data, output_path):
    """Writes an ACC-compatible Issues import .xlsx to output_path.

    issues_data: list of dicts (keyed by the snake_case names in COLUMNS,
    e.g. "title"/"due_date"), or a pandas DataFrame with matching column
    names -- either the snake_case keys or the exact header text. Pass an
    empty list/None to generate a headers-only template.
    """
    rows = _rows_from(issues_data)

    wb = Workbook()
    ws = wb.active
    ws.title = u"Issues"

    for col_idx, (header_text, _key, _align, _min_w, _max_w) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal=u"center", vertical=u"center", wrap_text=True)
        cell.border = GRID_BORDER
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    for row_offset, row in enumerate(rows):
        row_idx = row_offset + 2
        is_shaded = (row_offset % 2 == 1)
        for col_idx, (header_text, key, align, _min_w, _max_w) in enumerate(COLUMNS, start=1):
            value = _cell_value(row, header_text, key)
            cell = ws.cell(row=row_idx, column=col_idx, value=value if value != u"" else None)
            cell.alignment = Alignment(horizontal=align, vertical=u"center", wrap_text=(align == u"left"))
            cell.border = GRID_BORDER
            if is_shaded:
                cell.fill = ZEBRA_FILL
        ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT

    for col_idx, (header_text, key, _align, min_w, max_w) in enumerate(COLUMNS, start=1):
        longest = len(header_text)
        for row in rows:
            value = _cell_value(row, header_text, key)
            longest = max(longest, len(str(value)))
        width = max(min_w, min(max_w, longest + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = u"A2"

    wb.save(output_path)
    return output_path


def _load_issues_from_json(path):
    with open(path, u"r") as f:
        return json.load(f)


def main():
    parser = argparse.ArgumentParser(
        description=u"Generate an ACC-compatible Issues Excel import file.")
    parser.add_argument(
        u"--output", u"-o", default=u"ACC_Issues_Import_Template.xlsx",
        help=u"Output .xlsx path (default: %(default)s)")
    parser.add_argument(
        u"--input-json", u"-i", default=None,
        help=u"Optional path to a JSON file (array of issue objects) to populate the sheet. "
             u"Omit to generate a headers-only template.")
    args = parser.parse_args()

    issues_data = _load_issues_from_json(args.input_json) if args.input_json else []
    output_path = create_acc_issues_excel(issues_data, args.output)
    print(u"Wrote {} issue(s) to {}".format(len(issues_data), output_path))


if __name__ == u"__main__":
    main()
